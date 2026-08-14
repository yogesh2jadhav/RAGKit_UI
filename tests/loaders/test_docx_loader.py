"""
Purpose
-------
Tests the DOCX document loader.

Responsibilities
----------------
- Verify DOCX files are supported.
- Verify DOCX paragraphs are loaded.
- Verify DOCX tables are loaded.
- Verify empty table cells are ignored.
- Verify document metadata is populated.
- Verify source metadata is propagated.

Does NOT
--------
- Test chunking.
- Test embedding.
- Test vector storage.
"""

from pathlib import Path

from docx import Document as WordDocument

from ragkit.loaders.docx_loader import DocxLoader
from ragkit.models.source_document import SourceDocument


DOCX_MIME_TYPE = (
    "application/vnd.openxmlformats-officedocument"
    ".wordprocessingml.document"
)


def create_docx(
    path: Path,
    paragraphs: list[str],
) -> None:
    """
    Create a small DOCX file for testing.
    """

    document = WordDocument()

    for paragraph in paragraphs:
        document.add_paragraph(paragraph)

    document.save(path)


def create_docx_with_table(
    path: Path,
) -> None:
    """
    Create a DOCX containing a table.
    """

    document = WordDocument()

    table = document.add_table(
        rows=3,
        cols=2,
    )

    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Technology"

    table.cell(1, 0).text = "Apache Spark"
    table.cell(1, 1).text = "Distributed processing"

    table.cell(2, 0).text = "Delta Lake"
    table.cell(2, 1).text = "Data lakehouse"

    document.save(path)


def create_docx_with_empty_cells(
    path: Path,
) -> None:
    """
    Create a DOCX containing empty table cells.
    """

    document = WordDocument()

    table = document.add_table(
        rows=2,
        cols=3,
    )

    table.cell(0, 0).text = "Spark"
    table.cell(0, 1).text = ""
    table.cell(0, 2).text = "Databricks"

    table.cell(1, 0).text = ""
    table.cell(1, 1).text = "Delta Lake"
    table.cell(1, 2).text = ""

    document.save(path)


def test_docx_loader_supports_docx() -> None:
    """
    Verify DocxLoader supports .docx files.
    """

    source = SourceDocument(
        uri="example.docx",
        mime_type=DOCX_MIME_TYPE,
    )

    assert DocxLoader.supports(source) is True


def test_docx_loader_does_not_support_txt() -> None:
    """
    Verify DocxLoader does not support text files.
    """

    source = SourceDocument(
        uri="example.txt",
        mime_type="text/plain",
    )

    assert DocxLoader.supports(source) is False


def test_docx_loader_supports_uppercase_extension() -> None:
    """
    Verify extension matching is case-insensitive.
    """

    source = SourceDocument(
        uri="example.DOCX",
        mime_type=DOCX_MIME_TYPE,
    )

    assert DocxLoader.supports(source) is True


def test_docx_loader_loads_paragraphs(
    tmp_path: Path,
) -> None:
    """
    Verify paragraphs are extracted in document order.
    """

    path = tmp_path / "sample.docx"

    create_docx(
        path,
        [
            "Apache Spark is a distributed processing engine.",
            "Spark supports batch processing.",
            "Spark also supports streaming.",
        ],
    )

    source = SourceDocument(
        uri=str(path),
        mime_type=DOCX_MIME_TYPE,
    )

    document = DocxLoader().load(source)

    assert (
        document.content
        == (
            "Apache Spark is a distributed processing engine.\n"
            "Spark supports batch processing.\n"
            "Spark also supports streaming."
        )
    )


def test_docx_loader_loads_table_content(
    tmp_path: Path,
) -> None:
    """
    Verify table cells are extracted.
    """

    path = tmp_path / "table.docx"

    create_docx_with_table(path)

    source = SourceDocument(
        uri=str(path),
        mime_type=DOCX_MIME_TYPE,
    )

    document = DocxLoader().load(source)

    assert "Name | Technology" in document.content

    assert (
        "Apache Spark | Distributed processing"
        in document.content
    )

    assert (
        "Delta Lake | Data lakehouse"
        in document.content
    )


def test_docx_loader_ignores_empty_table_cells(
    tmp_path: Path,
) -> None:
    """
    Verify empty table cells do not create empty separators.
    """

    path = tmp_path / "empty_cells.docx"

    create_docx_with_empty_cells(path)

    source = SourceDocument(
        uri=str(path),
        mime_type=DOCX_MIME_TYPE,
    )

    document = DocxLoader().load(source)

    assert "Spark | Databricks" in document.content
    assert "Delta Lake" in document.content

    assert "Spark |  | Databricks" not in document.content


def test_docx_loader_populates_metadata(
    tmp_path: Path,
) -> None:
    """
    Verify standard document metadata is populated.
    """

    path = tmp_path / "spark_notes.docx"

    create_docx(
        path,
        [
            "Apache Spark notes.",
        ],
    )

    source = SourceDocument(
        uri=str(path),
        mime_type=DOCX_MIME_TYPE,
    )

    document = DocxLoader().load(source)

    assert document.metadata["uri"] == str(path)
    assert document.metadata["mime_type"] == DOCX_MIME_TYPE
    assert document.metadata["filename"] == "spark_notes.docx"


def test_docx_loader_propagates_source_metadata(
    tmp_path: Path,
) -> None:
    """
    Verify user-provided SourceDocument metadata is preserved.
    """

    path = tmp_path / "spark.docx"

    create_docx(
        path,
        [
            "Apache Spark.",
        ],
    )

    source = SourceDocument(
        uri=str(path),
        mime_type=DOCX_MIME_TYPE,
        metadata={
            "category": "spark",
            "level": "advanced",
        },
    )

    document = DocxLoader().load(source)

    assert document.metadata["category"] == "spark"
    assert document.metadata["level"] == "advanced"


def test_docx_loader_empty_document(
    tmp_path: Path,
) -> None:
    """
    Verify an empty DOCX document can be loaded.
    """

    path = tmp_path / "empty.docx"

    create_docx(
        path,
        [],
    )

    source = SourceDocument(
        uri=str(path),
        mime_type=DOCX_MIME_TYPE,
    )

    document = DocxLoader().load(source)

    assert document.content == ""